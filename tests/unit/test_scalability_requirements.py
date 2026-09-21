from __future__ import annotations

import ast
import os
import subprocess
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from construction_os.importers import parse_vor
from construction_os.references import ContractType, NoteType, SourceType, Unit, get_rate
from construction_os.storage.models import Base, CompanyRow, ObjectRow, ValueSourceRow
from construction_os.storage.repositories import (
    TENANT_REPOSITORIES,
    ImmutableRecordError,
    WorkItemRepository,
)

ROOT = Path(__file__).resolve().parents[2]
EXCLUDED_TABLES = {"companies", "reference_rates"}
FORBIDDEN_RATES = {0.20, 0.22, 1.22, 0.25, 0.30, 0.14}


def _shifted_fixture(source: Path, target: Path) -> None:
    workbook = load_workbook(source)
    sheet = workbook[workbook.sheetnames[0]]
    sheet.unmerge_cells("A40:E40")
    sheet.unmerge_cells("A41:F41")
    sheet.move_range("A9:F44", rows=2, cols=0, translate=True)
    sheet.merge_cells("A42:E42")
    sheet.merge_cells("A43:F43")
    workbook.save(target)


def test_M01_company_id_everywhere():
    for name, table in Base.metadata.tables.items():
        if name in EXCLUDED_TABLES:
            continue
        column = table.c.company_id
        assert column.nullable is False
        assert column.index is True or any(
            column.name in [c.name for c in index.columns] for index in table.indexes
        )


def test_M02_tenant_isolation_repository_catalog_complete():
    names = {repository.model.__tablename__ for repository in TENANT_REPOSITORIES}
    assert names == set(Base.metadata.tables) - EXCLUDED_TABLES


def test_M02_tenant_isolation_work_items(sqlite_session):
    company_a = CompanyRow(name="A")
    company_b = CompanyRow(name="B")
    sqlite_session.add_all([company_a, company_b])
    sqlite_session.flush()
    for company in (company_a, company_b):
        source = ValueSourceRow(company_id=company.id, source_type="document", confidence="exact")
        sqlite_session.add(source)
        sqlite_session.flush()
        obj = ObjectRow(company_id=company.id, name=f"O-{company.name}")
        sqlite_session.add(obj)
        sqlite_session.flush()
        WorkItemRepository(sqlite_session).add(
            company.id,
            object_id=obj.id,
            position_no=1,
            name="Работа",
            unit="шт",
            quantity=Decimal("1"),
            price_gross=Decimal("1"),
            amount_gross=Decimal("1"),
            vat_rate=Decimal("0.22"),
            source_id=source.id,
            valid_from=date(2026, 1, 1),
        )
    rows_a = WorkItemRepository(sqlite_session).list_current(company_a.id)
    rows_b = WorkItemRepository(sqlite_session).list_current(company_b.id)
    assert {row.company_id for row in rows_a} == {company_a.id}
    assert {row.company_id for row in rows_b} == {company_b.id}


def test_M03_values_are_superseded_not_updated(sqlite_session):
    company = CompanyRow(name="A")
    sqlite_session.add(company)
    sqlite_session.flush()
    source = ValueSourceRow(company_id=company.id, source_type="document", confidence="exact")
    sqlite_session.add(source)
    sqlite_session.flush()
    obj = ObjectRow(company_id=company.id, name="O")
    sqlite_session.add(obj)
    sqlite_session.flush()
    repo = WorkItemRepository(sqlite_session)
    old = repo.add(
        company.id,
        object_id=obj.id,
        position_no=1,
        name="Работа",
        unit="шт",
        quantity=Decimal("1"),
        price_gross=Decimal("10"),
        amount_gross=Decimal("10"),
        vat_rate=Decimal("0.22"),
        source_id=source.id,
        valid_from=date(2026, 1, 1),
    )
    new = repo.supersede(
        company.id, old.id, {"price_gross": Decimal("11")}, "correction", "test", date(2026, 2, 1)
    )
    assert new.id != old.id
    assert repo.get_on_date(company.id, old.id, date(2026, 1, 15)).price_gross == Decimal("10")
    with pytest.raises(ImmutableRecordError):
        repo.update(old.id)
    with pytest.raises(ImmutableRecordError):
        repo.delete(old.id)


def test_M04_no_rate_literals_outside_references():
    violations = []
    for path in (ROOT / "src").rglob("*.py"):
        if "references" in path.parts:
            continue
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if (
                isinstance(node, ast.Constant)
                and isinstance(node.value, float)
                and node.value in FORBIDDEN_RATES
            ):
                violations.append((path, node.lineno, node.value))
    assert violations == []


def test_M05_reference_catalog_is_mandatory():
    assert str(get_rate("VAT_RATE", date(2025, 12, 31)).value) == "0.20"
    assert str(get_rate("VAT_RATE", date(2026, 1, 1)).value) == "0.22"


def test_M06_domain_does_not_import_sqlalchemy():
    violations = []
    for path in (ROOT / "src" / "construction_os" / "domain").rglob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                violations.extend(
                    alias.name for alias in node.names if alias.name.startswith("sqlalchemy")
                )
            if isinstance(node, ast.ImportFrom) and (node.module or "").startswith("sqlalchemy"):
                violations.append(node.module)
    assert violations == []


def test_M07_calc_is_pure_from_storage_and_io():
    for path in (ROOT / "src" / "construction_os" / "calc").rglob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        imports = []
        for node in ast.walk(tree):
            if isinstance(node, ast.ImportFrom):
                imports.append(node.module or "")
            elif isinstance(node, ast.Import):
                imports.extend(alias.name for alias in node.names)
        assert not any(name.startswith("construction_os.storage") for name in imports)
        assert "openpyxl" not in imports


def test_M08_environment_settings_only():
    assert not (ROOT / ".env").exists()
    content = (ROOT / ".env.example").read_text(encoding="utf-8")
    assert content.strip() == "DATABASE_URL="


def test_M09_migrations_from_zero(tmp_path):
    database = tmp_path / "migration.db"
    env = os.environ.copy()
    env["DATABASE_URL"] = f"sqlite+pysqlite:///{database}"
    env["PYTHONPATH"] = str(ROOT / "src")
    for command in (["upgrade", "head"], ["downgrade", "base"], ["upgrade", "head"]):
        completed = subprocess.run(
            [sys.executable, "-m", "alembic", *command],
            cwd=ROOT,
            env=env,
            capture_output=True,
            text=True,
            check=False,
        )
        assert completed.returncode == 0, completed.stdout + completed.stderr


def test_M10_no_single_tenant_data_tables():
    data_tables = set(Base.metadata.tables) - EXCLUDED_TABLES
    assert data_tables
    assert all("company_id" in Base.metadata.tables[name].c for name in data_tables)


def test_M11_reference_types_are_extensible():
    assert Unit.PIECE.value == "шт"
    assert ContractType.UNKNOWN.value == "unknown"
    assert SourceType.DOCUMENT.value == "document"
    assert NoteType.RESOURCE_PLAN.value == "resource_plan"


def test_M11_calc_and_importers_do_not_branch_on_unit_values():
    files = list((ROOT / "src" / "construction_os" / "calc").rglob("*.py")) + list(
        (ROOT / "src" / "construction_os" / "importers").rglob("*.py")
    )
    text = "\n".join(path.read_text(encoding="utf-8") for path in files)
    assert "Unit.PIECE" not in text
    assert "Unit.CUBIC_METER" not in text


def test_M12_importer_finds_shifted_header(fixtures_dir, tmp_path):
    target = tmp_path / "shifted.xlsx"
    _shifted_fixture(fixtures_dir / "vor_object_a.xlsx", target)
    parsed = parse_vor(target)
    assert parsed.header_row == 11
    assert len(parsed.items) == 29
    assert parsed.total_gross == Decimal("35656922.00")


def test_M03_postgres_trigger_rejects_direct_update(db_session):
    if db_session.bind.dialect.name != "postgresql":
        pytest.skip("PostgreSQL trigger test")
    company = CompanyRow(name="Trigger Co")
    db_session.add(company)
    db_session.flush()
    source = ValueSourceRow(company_id=company.id, source_type="document", confidence="exact")
    db_session.add(source)
    db_session.flush()
    obj = ObjectRow(company_id=company.id, name="Trigger Object")
    db_session.add(obj)
    db_session.flush()
    item = WorkItemRepository(db_session).add(
        company.id,
        object_id=obj.id,
        position_no=1,
        name="Работа",
        unit="шт",
        quantity=Decimal("1"),
        price_gross=Decimal("10"),
        amount_gross=Decimal("10"),
        vat_rate=Decimal("0.22"),
        source_id=source.id,
        valid_from=date(2026, 1, 1),
    )
    from sqlalchemy import text

    with pytest.raises(Exception, match="immutable"):
        db_session.execute(
            text("UPDATE work_items SET price_gross = 11 WHERE id = :id"), {"id": item.id}
        )
