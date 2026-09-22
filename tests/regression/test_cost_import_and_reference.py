from datetime import date
from decimal import Decimal
from pathlib import Path
import subprocess,sys
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from construction_os.importers import CostImportError,article_codes,parse_costs,persist_costs
from construction_os.references import DEFAULT_COST_ARTICLES
from construction_os.storage.models import CompanyRow,CostArticleRow,CostEntryRow,ObjectRow,ValueRefRow
ROOT=Path(__file__).resolve().parents[2]
def seed_catalog(session):
    for i,(code,category,name) in enumerate(DEFAULT_COST_ARTICLES,1): session.add(CostArticleRow(code=code,category=category,name=name,is_active=True,sort_order=i))
    session.flush()
def make_book(tmp_path,rows):
    out=tmp_path/"costs.xlsx"; subprocess.run([sys.executable,"scripts/make_cost_template.py","--out",str(out)],cwd=ROOT,check=True)
    wb=load_workbook(out); ws=wb["Затраты"]
    for row in rows: ws.append(row)
    wb.save(out); return out
def test_template_has_required_columns(tmp_path):
    ws=load_workbook(make_book(tmp_path,[]),read_only=True)["Затраты"]; headers=[c.value for c in next(ws.iter_rows())]; assert "amount_type" in headers and "rate_value" in headers
def test_template_has_29_articles(tmp_path): assert load_workbook(make_book(tmp_path,[]),read_only=True)["Статьи"].max_row-1==29
def test_parse_fixed(tmp_path):
    p=parse_costs(make_book(tmp_path,[["O",None,"MAT",2,"шт",5,10,"fixed",None,"net","src",None,None]]),{x[0] for x in DEFAULT_COST_ARTICLES}); assert p.rows[0].amount==Decimal("10.00")
def test_parse_share(tmp_path):
    p=parse_costs(make_book(tmp_path,[["O",None,"FINANCE",None,None,None,None,"share_of_revenue",Decimal("0.027"),"net","src",None,None]]),{x[0] for x in DEFAULT_COST_ARTICLES}); assert p.rows[0].amount is None and p.rows[0].rate_value==Decimal("0.027")
def test_parse_unknown_article_has_row(tmp_path):
    path=make_book(tmp_path,[["O",None,"BAD",None,None,None,1,"fixed",None,"net",None,None,None]])
    with pytest.raises(CostImportError,match="строка 2"): parse_costs(path,{x[0] for x in DEFAULT_COST_ARTICLES})
@pytest.mark.parametrize("amount_type,amount,rate",[("fixed",None,None),("share_of_revenue",None,None),("fixed",1,Decimal("0.1")),("share_of_revenue",1,Decimal("0.1"))])
def test_parse_amount_shape_errors(tmp_path,amount_type,amount,rate):
    path=make_book(tmp_path,[["O",None,"MAT",None,None,None,amount,amount_type,rate,"net",None,None,None]])
    with pytest.raises(CostImportError): parse_costs(path,{x[0] for x in DEFAULT_COST_ARTICLES})
def test_parse_quantity_price_warning(tmp_path):
    p=parse_costs(make_book(tmp_path,[["O",None,"MAT",2,"шт",5,11,"fixed",None,"net",None,None,None]]),{x[0] for x in DEFAULT_COST_ARTICLES}); assert "quantity × price" in p.warnings[0]
def test_persist_and_idempotency(sqlite_session,tmp_path):
    seed_catalog(sqlite_session); c=CompanyRow(name="A"); sqlite_session.add(c); sqlite_session.flush(); obj=ObjectRow(company_id=c.id,name="O",valid_from=date(2026,1,1)); sqlite_session.add(obj); sqlite_session.flush()
    path=make_book(tmp_path,[["O",None,"MAT",None,None,None,10,"fixed",None,"net","src",None,None]]); p=parse_costs(path,article_codes(sqlite_session)); first=persist_costs(sqlite_session,"A",p,path,date(2026,1,1)); sqlite_session.commit(); second=persist_costs(sqlite_session,"A",p,path,date(2026,1,1)); assert not first.skipped_duplicate and second.skipped_duplicate; assert len(sqlite_session.scalars(select(CostEntryRow)).all())==1
def test_persist_provenance(sqlite_session,tmp_path):
    seed_catalog(sqlite_session); c=CompanyRow(name="A"); sqlite_session.add(c); sqlite_session.flush(); sqlite_session.add(ObjectRow(company_id=c.id,name="O",valid_from=date(2026,1,1))); sqlite_session.flush(); path=make_book(tmp_path,[["O",None,"MAT",None,None,None,10,"fixed",None,"net",None,None,None]]); p=parse_costs(path,article_codes(sqlite_session)); persist_costs(sqlite_session,"A",p,path,date(2026,1,1)); ref=sqlite_session.scalar(select(ValueRefRow).where(ValueRefRow.entity_name=="cost_entries")); assert ref.field_name=="amount"
def test_reference_values_file_matches_generator():
    completed=subprocess.run([sys.executable,"scripts/make_reference_values.py"],cwd=ROOT,capture_output=True,text=True,check=True); assert completed.stdout==(ROOT/"docs/spec/reference_values.md").read_text(encoding="utf-8")
