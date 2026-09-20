from decimal import Decimal
from openpyxl import Workbook
from construction_os.importers import parse_schedule

def test_schedule_source_data_and_periods(tmp_path):
    p=tmp_path/"s.xlsx";wb=Workbook();ws=wb.active;ws.title="Календарный график"
    ws["A2"]="Объект. Период работ: 05.10.2026–15.11.2026."
    ws["A4"]="Ресурсный план: 05–09.10 до 10 чел.; далее Дятлинка и Ямуга делят мобильную бригаду 10 чел."
    ws["A5"]="Реверс: сторона 1 → сторона 2"
    headers=["№ сметы","Наименование работ","Фронт/сторона","Ед.","Объем","Начало","Окончание","Дни","Звено, чел.","Стоимость, руб.","Комментарий"]
    for c,v in enumerate(headers,1):ws.cell(7,c,v)
    ws.cell(9,1,1);ws.cell(9,2,"Работа");ws.cell(9,3,"весь объект");ws.cell(9,4,"шт");ws.cell(9,5,2);ws.cell(9,6,"05.10.2026");ws.cell(9,7,"05.10.2026");ws.cell(9,8,1);ws.cell(9,9,6);ws.cell(9,10,Decimal("85720.38"));ws.cell(9,13,2)
    wb.save(p);s=parse_schedule(p)
    assert s.tasks[0].days==1 and s.tasks[0].crew_size==Decimal("6")
    assert s.period_mismatches==() and s.total_amount==Decimal("85720.38")
    resource=[n for n in s.notes if n["type"]=="resource_plan"][0]["shared_resource"]
    assert resource["from"]=="2026-10-10" and resource["crew"]=="10"
