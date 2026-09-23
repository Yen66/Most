from __future__ import annotations

import json
import re
from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from construction_os.money import as_decimal, money, round_position, sum_positions


class VorImportError(ValueError):
    """Raised when a VOR workbook cannot be parsed deterministically."""


@dataclass(frozen=True, slots=True)
class Item:
    position_no: int
    name: str
    unit: str
    quantity: Decimal
    price_gross: Decimal
    amount_gross: Decimal
    row_no: int
    quantity_cell: str
    price_cell: str
    amount_cell: str
    raw_price: str
    amount_formula: str | None


@dataclass(frozen=True, slots=True)
class ParsedVor:
    items: tuple[Item, ...]
    total_gross: Decimal
    vat_rate: Decimal
    price_is_final: bool
    discrepancies: tuple[str, ...]
    sha256: str
    header_row: int
    sheet_name: str


def _header(sheet) -> int:
    for row in range(1, sheet.max_row + 1):
        text = " | ".join(
            str(sheet.cell(row, column).value or "").lower()
            for column in range(1, min(sheet.max_column, 8) + 1)
        )
        if all(token in text for token in ("наименование", "количество", "расцен", "ндс")):
            return row
    raise VorImportError("header not found")


def _vat_and_total_row(sheet) -> tuple[Decimal, int]:
    pattern = re.compile(r"ндс\s*(\d+(?:[.,]\d+)?)\s*%", re.IGNORECASE)
    for row in range(1, sheet.max_row + 1):
        for column in range(1, min(sheet.max_column, 8) + 1):
            value = sheet.cell(row, column).value
            if isinstance(value, str) and (match := pattern.search(value)):
                return Decimal(match.group(1).replace(",", ".")) / Decimal("100"), row
    raise VorImportError("VAT not stated")


def _position_number(value, formula_value, previous: int | None) -> int | None:
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(formula_value, str) and formula_value.startswith("=") and previous is not None:
        return previous + 1
    return None


def parse_vor(path: str | Path) -> ParsedVor:
    source_path = Path(path)
    formulas_book = load_workbook(source_path, data_only=False)
    values_book = load_workbook(source_path, data_only=True)
    formulas_sheet = formulas_book[formulas_book.sheetnames[0]]
    values_sheet = values_book[formulas_sheet.title]
    header_row = _header(formulas_sheet)
    vat_rate, total_row = _vat_and_total_row(formulas_sheet)
    items: list[Item] = []
    discrepancies: list[str] = []
    previous_position: int | None = None
    for row in range(header_row + 1, total_row):
        position = _position_number(
            values_sheet.cell(row, 1).value,
            formulas_sheet.cell(row, 1).value,
            previous_position,
        )
        name = values_sheet.cell(row, 2).value
        if position is None or not isinstance(name, str) or not name.strip():
            continue
        quantity = as_decimal(values_sheet.cell(row, 4).value)
        raw_price = values_sheet.cell(row, 5).value
        price = money(raw_price)
        expected = round_position(quantity, price)
        cached_amount = values_sheet.cell(row, 6).value
        observed = money(cached_amount) if cached_amount is not None else expected
        if observed != expected:
            discrepancies.append(f"row {row}: {observed} != {expected}")
        items.append(
            Item(
                position_no=position,
                name=name.strip(),
                unit=str(values_sheet.cell(row, 3).value or "").strip(),
                quantity=quantity,
                price_gross=price,
                amount_gross=observed,
                row_no=row,
                quantity_cell=f"D{row}",
                price_cell=f"E{row}",
                amount_cell=f"F{row}",
                raw_price=str(raw_price),
                amount_formula=formulas_sheet.cell(row, 6).value,
            )
        )
        previous_position = position
    total = sum_positions(item.amount_gross for item in items)
    cached_total = values_sheet.cell(total_row, 6).value
    if cached_total is not None and money(cached_total) != total:
        raise VorImportError(f"total {cached_total} != {total}")
    all_text = " ".join(
        str(formulas_sheet.cell(row, column).value or "")
        for row in range(1, formulas_sheet.max_row + 1)
        for column in range(1, min(formulas_sheet.max_column, 7) + 1)
    ).lower()
    price_is_final = "скоррект" not in all_text
    semantic_payload = {
        "sheet": formulas_sheet.title,
        "vat_rate": str(vat_rate),
        "price_is_final": price_is_final,
        "items": [
            [
                item.position_no,
                item.name,
                item.unit,
                str(item.quantity),
                str(item.price_gross),
                str(item.amount_gross),
            ]
            for item in items
        ],
    }
    content_sha256 = sha256(
        json.dumps(semantic_payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()
    return ParsedVor(
        items=tuple(items),
        total_gross=total,
        vat_rate=vat_rate,
        price_is_final=price_is_final,
        discrepancies=tuple(discrepancies),
        sha256=content_sha256,
        header_row=header_row,
        sheet_name=formulas_sheet.title,
    )
