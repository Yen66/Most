from pathlib import Path

from openpyxl import Workbook

from construction_os.importers.costs import HEADERS
from construction_os.references import DEFAULT_COST_ARTICLES


def make_template(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Затраты"
    sheet.append(list(HEADERS))
    catalog = workbook.create_sheet("Статьи")
    catalog.append(["code", "category", "name"])
    for row in DEFAULT_COST_ARTICLES:
        catalog.append(row)
    info = workbook.create_sheet("Инструкция")
    info.append(["Для fixed заполните amount; для share_of_revenue amount оставьте пустым и заполните rate_value."])
    info.append(["Пустой amount_type означает fixed. vat_mode: gross / net / unknown."])
    workbook.save(path)
