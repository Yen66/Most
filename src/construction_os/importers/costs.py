from __future__ import annotations
import hashlib
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from construction_os.money import as_decimal, money

class CostImportError(ValueError):
    """Invalid cost workbook."""

@dataclass(frozen=True, slots=True)
class ParsedCost:
    row_no:int; object_name:str; work_position_no:int|None; article_code:str
    quantity:Decimal|None; unit:str|None; price:Decimal|None; amount:Decimal|None
    amount_type:str; rate_value:Decimal|None; vat_mode:str; source:str|None
    source_date:date|None; note:str|None

@dataclass(frozen=True, slots=True)
class ParsedCosts:
    rows:tuple[ParsedCost,...]; sheet_name:str; sha256:str; warnings:tuple[str,...]

HEADERS=("object","work_position_no","article_code","quantity","unit","price","amount","amount_type","rate_value","vat_mode","source","source_date","note")

def _decimal(value):
    return None if value in (None,"") else as_decimal(value)

def parse_costs(path: str|Path, article_codes:set[str]) -> ParsedCosts:
    source_path=Path(path); digest=hashlib.sha256(source_path.read_bytes()).hexdigest()
    workbook=load_workbook(source_path,data_only=True,read_only=True); sheet=workbook[workbook.sheetnames[0]]
    headers=[str(c.value or "").strip() for c in next(sheet.iter_rows(min_row=1,max_row=1))]
    if headers[:len(HEADERS)]!=list(HEADERS): raise CostImportError("строка 1: неверные колонки шаблона затрат")
    rows=[]; warnings=[]
    for row_no,cells in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
        values=dict(zip(HEADERS,cells,strict=False))
        if all(values.get(n) in (None,"") for n in HEADERS): continue
        object_name=str(values.get("object") or "").strip(); code=str(values.get("article_code") or "").strip()
        amount_type=str(values.get("amount_type") or "fixed").strip(); vat_mode=str(values.get("vat_mode") or "unknown").strip()
        if not object_name: raise CostImportError(f"строка {row_no}: object обязателен")
        if code not in article_codes: raise CostImportError(f"строка {row_no}: неизвестный article_code {code!r}")
        if amount_type not in {"fixed","share_of_revenue"}: raise CostImportError(f"строка {row_no}: неизвестный amount_type {amount_type!r}")
        if vat_mode not in {"gross","net","unknown"}: raise CostImportError(f"строка {row_no}: неизвестный vat_mode {vat_mode!r}")
        amount=_decimal(values.get("amount")); rate=_decimal(values.get("rate_value"))
        if (amount is None)==(rate is None): raise CostImportError(f"строка {row_no}: заполните ровно одно из amount или rate_value")
        if amount_type=="fixed" and amount is None: raise CostImportError(f"строка {row_no}: fixed требует amount")
        if amount_type=="share_of_revenue" and rate is None: raise CostImportError(f"строка {row_no}: share_of_revenue требует rate_value")
        quantity=_decimal(values.get("quantity")); price=_decimal(values.get("price"))
        if amount_type=="fixed" and quantity is not None and price is not None:
            expected=money(quantity*price)
            if money(amount)!=expected: warnings.append(f"строка {row_no}: amount {money(amount)} не равен quantity × price {expected}")
        pos=values.get("work_position_no")
        rows.append(ParsedCost(row_no,object_name,int(pos) if pos not in (None,"") else None,code,quantity,str(values["unit"]).strip() if values.get("unit") else None,price,money(amount) if amount is not None else None,amount_type,rate,vat_mode,str(values["source"]).strip() if values.get("source") else None,values["source_date"] if isinstance(values.get("source_date"),date) else None,str(values["note"]).strip() if values.get("note") else None))
    return ParsedCosts(tuple(rows),sheet.title,digest,tuple(warnings))
