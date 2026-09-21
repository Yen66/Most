from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from construction_os.money import as_decimal, money, sum_positions

from .vor import ParsedVor


@dataclass(frozen=True, slots=True)
class Task:
    position_no: int | None
    name: str
    front: str | None
    unit: str | None
    quantity: Decimal | None
    start_on: date | None
    end_on: date | None
    days: int | None
    crew_size: Decimal | None
    amount: Decimal | None
    period_volumes: tuple[Decimal | None, ...]
    row_no: int


@dataclass(frozen=True, slots=True)
class ParsedSchedule:
    tasks: tuple[Task, ...]
    notes: tuple[dict, ...]
    total_amount: Decimal
    period_mismatches: tuple[str, ...]
    sha256: str
    sheet_name: str
    object_name: str | None


def _date_value(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value), "%d.%m.%Y").date()


def _decimal_or_none(value) -> Decimal | None:
    return None if value is None else as_decimal(value)


def _find_header(sheet) -> int:
    for row in range(1, sheet.max_row + 1):
        text = " | ".join(
            str(sheet.cell(row, column).value or "").lower()
            for column in range(1, min(sheet.max_column, 12) + 1)
        )
        if all(token in text for token in ("наименование работ", "начало", "окончание", "звено")):
            return row
    raise ValueError("schedule header not found")


def _parse_notes(sheet) -> tuple[list[dict], str | None]:
    notes: list[dict] = []
    year: int | None = None
    object_name: str | None = None
    title = sheet["A1"].value
    if isinstance(title, str) and "—" in title:
        object_name = title.rsplit("—", 1)[-1].strip().title()
    for row in range(1, sheet.max_row + 1):
        text = sheet.cell(row, 1).value
        if not isinstance(text, str):
            continue
        lower = text.lower()
        if "период работ:" in lower:
            match = re.search(r"(\d{2}\.\d{2}\.\d{4})[–-](\d{2}\.\d{2}\.\d{4})", text)
            if match:
                year = int(match.group(2)[-4:])
                notes.append(
                    {
                        "type": "period",
                        "cell": f"A{row}",
                        "start": match.group(1),
                        "end": match.group(2),
                        "text": text,
                    }
                )
        elif lower.startswith("ресурсный план:"):
            note = {"type": "resource_plan", "cell": f"A{row}", "text": text}
            match = re.search(
                r"\d{2}[–-](\d{2}\.\d{2}).*?далее\s+(.+?)\s+и\s+(.+?)\s+делят\s+мобильную\s+бригаду\s+(\d+)\s+чел",
                lower,
            )
            if match and year:
                day, month = map(int, match.group(1).split("."))
                note["shared_resource"] = {
                    "resource": "мобильная бригада",
                    "objects": [match.group(2).title(), match.group(3).title()],
                    "from": (date(year, month, day) + timedelta(days=1)).isoformat(),
                    "crew": match.group(4),
                }
            notes.append(note)
        elif lower.startswith("реверс:"):
            notes.append({"type": "reverse_scheme", "cell": f"A{row}", "text": text})
        elif lower.startswith("примечание:"):
            notes.append({"type": "work_regime", "cell": f"A{row}", "text": text})
    return notes, object_name


def parse_schedule(path: str | Path) -> ParsedSchedule:
    source_path = Path(path)
    workbook = load_workbook(source_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header_row = _find_header(sheet)
    notes, object_name = _parse_notes(sheet)
    tasks: list[Task] = []
    period_mismatches: list[str] = []
    for row in range(header_row + 1, sheet.max_row + 1):
        name = sheet.cell(row, 2).value
        if not isinstance(name, str) or not name.strip():
            continue
        position_value = sheet.cell(row, 1).value
        position = int(position_value) if isinstance(position_value, (int, float)) else None
        quantity = _decimal_or_none(sheet.cell(row, 5).value)
        periods = tuple(_decimal_or_none(sheet.cell(row, column).value) for column in range(12, 24))
        if quantity is not None:
            period_total = sum((value for value in periods if value is not None), Decimal("0"))
            if period_total != quantity:
                period_mismatches.append(f"row {row}: {period_total} != {quantity}")
        amount_value = sheet.cell(row, 10).value
        tasks.append(
            Task(
                position_no=position,
                name=name.strip(),
                front=sheet.cell(row, 3).value,
                unit=sheet.cell(row, 4).value,
                quantity=quantity,
                start_on=_date_value(sheet.cell(row, 6).value),
                end_on=_date_value(sheet.cell(row, 7).value),
                days=sheet.cell(row, 8).value,
                crew_size=_decimal_or_none(sheet.cell(row, 9).value),
                amount=money(amount_value) if amount_value is not None else None,
                period_volumes=periods,
                row_no=row,
            )
        )
    return ParsedSchedule(
        tasks=tuple(tasks),
        notes=tuple(notes),
        total_amount=sum_positions(task.amount for task in tasks if task.amount is not None),
        period_mismatches=tuple(period_mismatches),
        sha256=sha256(source_path.read_bytes()).hexdigest(),
        sheet_name=sheet.title,
        object_name=object_name,
    )


def reconcile(vor: ParsedVor, schedule: ParsedSchedule) -> list[tuple[int, str]]:
    quantities: dict[int, Decimal] = {}
    amounts: dict[int, Decimal] = {}
    for task in schedule.tasks:
        if task.position_no is None:
            continue
        quantities[task.position_no] = quantities.get(task.position_no, Decimal("0")) + (
            task.quantity or Decimal("0")
        )
        amounts[task.position_no] = amounts.get(task.position_no, Decimal("0")) + (
            task.amount or Decimal("0")
        )
    differences: list[tuple[int, str]] = []
    for item in vor.items:
        if quantities.get(item.position_no, Decimal("0")) != item.quantity:
            differences.append((item.position_no, "quantity"))
        if money(amounts.get(item.position_no, Decimal("0"))) != item.amount_gross:
            differences.append((item.position_no, "amount"))
    return differences
