from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path
from zipfile import is_zipfile

from openpyxl import load_workbook
from sqlalchemy import select

from construction_os.importers.cost_persist import article_codes, persist_costs
from construction_os.importers.costs import parse_costs
from construction_os.importers.persist import persist_schedule, persist_vor
from construction_os.importers.schedule import parse_schedule
from construction_os.importers.vor import parse_vor
from construction_os.storage.models import CompanyRow, ObjectRow
from construction_os.storage.queries import object_revenues, portfolio_revenue


@dataclass
class ReceiptItem:
    file: str
    sha256: str | None = None
    type: str = "unknown"
    action: str = "rejected"
    positions: int = 0
    total_gross: str | None = None
    objects: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _classify(path: Path, codes: set[str]):
    """Use workbook structure; filenames never determine the parser."""
    suffix = path.suffix.lower()
    if suffix == ".xml":
        return "xml", None, ["XML: предположительно смета GGE — парсер не реализован"]
    if suffix in {".docx", ".pdf"}:
        return suffix[1:], None, ["неподдерживаемый формат: договоры и сканы — задачи T1/T2"]
    if not path.is_file():
        return "unknown", None, ["файл не найден"]
    if not is_zipfile(path):
        return "unknown", None, ["формат не определён — нет данных"]
    # openpyxl requires an xlsx extension for paths; a file handle permits a renamed workbook.
    try:
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=True)
            names = book.sheetnames
            book.close()
    except Exception as exc:
        return "xlsx", None, [f"повреждённый xlsx: {exc}"]
    with tempfile.TemporaryDirectory(prefix="construction-intake-") as temp:
        candidate = Path(temp) / "document.xlsx"
        shutil.copyfile(path, candidate)
        for kind, parser in (("vor", parse_vor), ("schedule", parse_schedule)):
            try:
                parsed = parser(candidate)
                if (kind == "vor" and parsed.items) or (kind == "schedule" and parsed.tasks):
                    return kind, parsed, []
            except (ValueError, TypeError, KeyError, OSError, IndexError):
                pass
        if "Статьи" in names:
            try:
                parsed = parse_costs(candidate, codes)
                return "costs", parsed, []
            except (ValueError, TypeError, KeyError, OSError, IndexError) as exc:
                return "costs", None, [str(exc)]
    return "xlsx", None, ["структура xlsx не распознана — нет данных"]


def _money_string(value: Decimal | None) -> str | None:
    return f"{value:.2f}" if value is not None else None


def intake_batch(session, company_name: str, paths: list[str | Path], on_date: date | None = None):
    if not company_name or not company_name.strip():
        raise ValueError("укажите --company")
    on_date = on_date or date.today()
    codes = article_codes(session)
    receipts = []
    detected = []
    for path_arg in paths:
        path = Path(path_arg)
        item = ReceiptItem(file=str(path))
        if path.is_file():
            item.sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
        kind, parsed, errors = _classify(path, codes)
        item.type = kind
        item.errors.extend(errors)
        if parsed is not None:
            if kind == "vor":
                item.positions = len(parsed.items)
                item.total_gross = _money_string(parsed.total_gross)
                item.warnings.extend(parsed.discrepancies)
            elif kind == "schedule":
                item.positions = len(parsed.tasks)
                item.total_gross = _money_string(parsed.total_amount)
                item.warnings.extend(parsed.period_mismatches)
            else:
                item.positions = len(parsed.rows)
                item.objects = sorted({row.object_name for row in parsed.rows})
                item.warnings.extend(parsed.warnings)
        receipts.append(item)
        detected.append((path, kind, parsed, item))
    # Dependencies: a cost sheet can appear before the VOR in the argument list.
    for kind_to_import in ("vor", "schedule", "costs"):
        for path, kind, parsed, item in detected:
            if kind != kind_to_import or parsed is None:
                continue
            try:
                with session.begin_nested():
                    if kind == "vor":
                        result = persist_vor(session, company_name, parsed, path, imported_on=on_date)
                    elif kind == "schedule":
                        result = persist_schedule(session, company_name, parsed, path)
                    else:
                        result = persist_costs(session, company_name, parsed, path, imported_on=on_date)
                    item.action = "skipped" if result.skipped_duplicate else "imported"
                    if result.object_id is not None:
                        name = session.scalar(
                            select(ObjectRow.name).where(ObjectRow.id == result.object_id)
                        )
                        if name and name not in item.objects:
                            item.objects.append(name)
                    if result.skipped_duplicate:
                        item.warnings.append("skipped (sha256)")
            except Exception as exc:
                item.action = "rejected"
                item.errors.append(str(exc))
    session.flush()
    rows = object_revenues(session, on_date, company_name)
    if rows:
        portfolio = portfolio_revenue(rows, on_date)
        gross, net = portfolio.gross, portfolio.net
    else:
        gross = net = Decimal("0")
    return {
        "files": [asdict(item) for item in receipts],
        "summary": {
            "imported": sum(item.action == "imported" for item in receipts),
            "skipped": sum(item.action == "skipped" for item in receipts),
            "rejected": sum(item.action == "rejected" for item in receipts),
            "portfolio_gross": _money_string(gross),
            "portfolio_net": _money_string(net),
        },
    }


def receipt_json(receipt: dict) -> str:
    return json.dumps(receipt, ensure_ascii=False, indent=2)
