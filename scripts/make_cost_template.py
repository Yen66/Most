from __future__ import annotations
import argparse
from pathlib import Path
from openpyxl import Workbook
from construction_os.importers.costs import HEADERS
from construction_os.references import DEFAULT_COST_ARTICLES


def make_template(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Затраты"
    ws.append(list(HEADERS))
    catalog = wb.create_sheet("Статьи")
    catalog.append(["code", "category", "name"])
    for row in DEFAULT_COST_ARTICLES:
        catalog.append(row)
    info = wb.create_sheet("Инструкция")
    info.append(
        [
            "Для fixed заполните amount; для share_of_revenue amount оставьте пустым и заполните rate_value."
        ]
    )
    info.append(["Пустой amount_type означает fixed. vat_mode: gross / net / unknown."])
    wb.save(path)


def main(argv=None) -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--out", required=True)
    a = p.parse_args(argv)
    make_template(Path(a.out))
    print(f"cost template: {a.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
